"""
Document Service for AgentOS.

Responsibilities:
- Create DOCX documents.
- Create PDF documents.
- Add titles and headings.
- Add paragraphs.
- Add bullet points.
- Add tables.
- Save generated documents locally.

Important:
- This service does NOT call the LLM.
- This service does NOT perform research.
- This service does NOT process the complete dataset.
- It only converts already-prepared content into documents.
"""

import logging
from pathlib import Path
from typing import Any, Dict, List, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from config.settings import get_settings


logger = logging.getLogger(__name__)

settings = get_settings()


class DocumentService:
    """
    Handles creation of DOCX and PDF documents.

    This class is intentionally independent from the LLM.
    """

    # ============================================================
    # Initialization
    # ============================================================

    def __init__(
        self,
        output_directory: Optional[str] = None,
    ):
        """
        Create the document output directory.
        """

        self.output_directory = Path(
            output_directory or settings.REPORTS_DIR
        )

        self.output_directory.mkdir(
            parents=True,
            exist_ok=True,
        )

    # ============================================================
    # Generate file path
    # ============================================================

    def get_output_path(
        self,
        filename: str,
    ) -> Path:
        """
        Return a safe output path inside the
        generated documents directory.
        """

        filename = Path(
            filename
        ).name

        return (
            self.output_directory
            / filename
        )

    # ============================================================
    # Create DOCX
    # ============================================================

    def create_docx(
        self,
        title: str,
        sections: List[Dict[str, Any]],
        filename: str,
    ) -> str:
        """
        Create a DOCX document.

        Returns:
            Path to the generated DOCX file.
        """

        output_path = self.get_output_path(
            filename
        )

        if output_path.suffix.lower() != ".docx":

            output_path = output_path.with_suffix(
                ".docx"
            )

        logger.info(
            "Creating DOCX document | path=%s",
            output_path,
        )

        document = Document()

        # --------------------------------------------------------
        # Title
        # --------------------------------------------------------

        title_paragraph = document.add_paragraph()

        title_paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        title_run = (
            title_paragraph.add_run(
                str(title)
            )
        )

        title_run.bold = True
        title_run.font.size = Pt(20)

        # --------------------------------------------------------
        # Sections
        # --------------------------------------------------------

        for section in sections:

            heading = section.get(
                "heading"
            )

            if heading:

                document.add_heading(
                    str(heading),
                    level=1,
                )

            paragraphs = section.get(
                "paragraphs",
                [],
            )

            for text in paragraphs:

                document.add_paragraph(
                    str(text)
                )

            bullets = section.get(
                "bullets",
                [],
            )

            for bullet in bullets:

                document.add_paragraph(
                    str(bullet),
                    style="List Bullet",
                )

            table_data = section.get(
                "table"
            )

            if table_data:

                self._add_docx_table(
                    document,
                    table_data,
                )

        document.save(
            output_path
        )

        logger.info(
            "DOCX document created successfully | path=%s",
            output_path,
        )

        return str(output_path)

    # ============================================================
    # DOCX Table
    # ============================================================

    def _add_docx_table(
        self,
        document: Document,
        table_data: List[List[Any]],
    ) -> None:
        """
        Add a table to a DOCX document.

        The first row is treated as the header.
        """

        if not table_data:
            return

        column_count = len(
            table_data[0]
        )

        table = document.add_table(
            rows=1,
            cols=column_count,
        )

        table.style = "Table Grid"

        header_cells = table.rows[0].cells

        for index, value in enumerate(
            table_data[0]
        ):

            header_cells[index].text = str(
                value
            )

        for row in table_data[1:]:

            cells = table.add_row().cells

            for index, value in enumerate(row):

                cells[index].text = str(
                    value
                )

    # ============================================================
    # Create PDF
    # ============================================================

    def create_pdf(
        self,
        title: str,
        sections: List[Dict[str, Any]],
        filename: str,
    ) -> str:
        """
        Create a PDF document.

        Returns:
            Path to the generated PDF file.
        """

        output_path = self.get_output_path(
            filename
        )

        if output_path.suffix.lower() != ".pdf":

            output_path = output_path.with_suffix(
                ".pdf"
            )

        logger.info(
            "Creating PDF document | path=%s",
            output_path,
        )

        styles = getSampleStyleSheet()

        title_style = styles["Title"]
        heading_style = styles["Heading1"]
        body_style = styles["BodyText"]

        document = SimpleDocTemplate(
            str(output_path),
            pagesize=A4,
            rightMargin=0.6 * inch,
            leftMargin=0.6 * inch,
            topMargin=0.6 * inch,
            bottomMargin=0.6 * inch,
        )

        elements = []

        # --------------------------------------------------------
        # Title
        # --------------------------------------------------------

        elements.append(
            Paragraph(
                str(title),
                title_style,
            )
        )

        elements.append(
            Spacer(
                1,
                0.25 * inch,
            )
        )

        # --------------------------------------------------------
        # Sections
        # --------------------------------------------------------

        for section in sections:

            heading = section.get(
                "heading"
            )

            if heading:

                elements.append(
                    Paragraph(
                        str(heading),
                        heading_style,
                    )
                )

                elements.append(
                    Spacer(
                        1,
                        0.08 * inch,
                    )
                )

            paragraphs = section.get(
                "paragraphs",
                [],
            )

            for text in paragraphs:

                elements.append(
                    Paragraph(
                        str(text),
                        body_style,
                    )
                )

                elements.append(
                    Spacer(
                        1,
                        0.1 * inch,
                    )
                )

            bullets = section.get(
                "bullets",
                [],
            )

            for bullet in bullets:

                elements.append(
                    Paragraph(
                        f"• {bullet}",
                        body_style,
                    )
                )

                elements.append(
                    Spacer(
                        1,
                        0.05 * inch,
                    )
                )

            table_data = section.get(
                "table"
            )

            if table_data:

                self._add_pdf_table(
                    elements,
                    table_data,
                )

                elements.append(
                    Spacer(
                        1,
                        0.15 * inch,
                    )
                )

        document.build(
            elements
        )

        logger.info(
            "PDF document created successfully | path=%s",
            output_path,
        )

        return str(output_path)

    # ============================================================
    # PDF Table
    # ============================================================

    def _add_pdf_table(
        self,
        elements: List[Any],
        table_data: List[List[Any]],
    ) -> None:
        """
        Add a table to a PDF.
        """

        if not table_data:
            return

        formatted_data = []

        for row in table_data:

            formatted_data.append(
                [
                    str(value)
                    for value in row
                ]
            )

        table = Table(
            formatted_data,
            repeatRows=1,
        )

        table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.lightgrey,
                    ),
                    (
                        "TEXTCOLOR",
                        (0, 0),
                        (-1, 0),
                        colors.black,
                    ),
                    (
                        "FONTNAME",
                        (0, 0),
                        (-1, 0),
                        "Helvetica-Bold",
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.5,
                        colors.grey,
                    ),
                    (
                        "VALIGN",
                        (0, 0),
                        (-1, -1),
                        "TOP",
                    ),
                    (
                        "LEFTPADDING",
                        (0, 0),
                        (-1, -1),
                        6,
                    ),
                    (
                        "RIGHTPADDING",
                        (0, 0),
                        (-1, -1),
                        6,
                    ),
                ]
            )
        )

        elements.append(
            table
        )


# ================================================================
# Singleton
# ================================================================

document_service = DocumentService()