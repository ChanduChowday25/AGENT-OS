"""
Document ingestion pipeline for AgentOS.

Supported document formats:

- PDF
- DOCX
- TXT

Pipeline:

File
↓
Extract text with metadata
↓
Normalize text
↓
Chunk text
↓
Generate embeddings
↓
Store in ChromaDB with source/page metadata
"""

import hashlib
import logging
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from docx import Document
from pypdf import PdfReader

from config.settings import get_settings
from database.chroma_client import get_document_collection
from services.embedding_service import embedding_service


logger = logging.getLogger(__name__)

settings = get_settings()


class DocumentIngestionService:
    """
    Converts uploaded documents into searchable ChromaDB chunks.

    Important:
    Page/source metadata is preserved during chunk creation so
    that the RAG Agent can provide reliable citations later.
    """

    def __init__(self):
        self.collection = get_document_collection()

    # ============================================================
    # Public API
    # ============================================================

    def ingest(
        self,
        file_id: str,
        file_path: str,
    ) -> Dict:
        """
        Ingest one document into ChromaDB.

        Returns ingestion metadata.
        """

        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(
                f"Document not found: {file_path}"
            )

        extension = path.suffix.lower()

        if extension not in settings.SUPPORTED_DOCUMENT_TYPES:
            raise ValueError(
                f"Unsupported document type: {extension}"
            )

        logger.info(
            "Starting ingestion: file_id=%s, type=%s",
            file_id,
            extension,
        )

        # --------------------------------------------------------
        # 1. Extract text while preserving page information
        # --------------------------------------------------------

        page_contents = self._extract_document_pages(
            path,
            extension,
        )

        if not page_contents:
            raise ValueError(
                "No readable text was extracted from the document."
            )

        # --------------------------------------------------------
        # 2. Normalize and chunk while preserving page metadata
        # --------------------------------------------------------

        chunk_records = []

        for page_number, page_text in page_contents:

            normalized_text = self._normalize_text(
                page_text
            )

            if not normalized_text:
                continue

            page_chunks = self._chunk_text(
                normalized_text,
                chunk_size=settings.CHUNK_SIZE,
                chunk_overlap=settings.CHUNK_OVERLAP,
            )

            for chunk in page_chunks:
                chunk_records.append(
                    {
                        "text": chunk,
                        "page": page_number,
                    }
                )

        if not chunk_records:
            raise ValueError(
                "Document produced no usable chunks."
            )

        chunks = [
            record["text"]
            for record in chunk_records
        ]

        logger.info(
            "Created %d chunks for file_id=%s",
            len(chunks),
            file_id,
        )

        # --------------------------------------------------------
        # 3. Generate embeddings
        # --------------------------------------------------------

        embeddings = embedding_service.embed_documents(
            chunks
        )

        if len(embeddings) != len(chunks):
            raise RuntimeError(
                "Embedding count does not match chunk count."
            )

        # --------------------------------------------------------
        # 4. Remove previous ingestion of same file
        # --------------------------------------------------------

        self.collection.delete(
            where={
                "file_id": file_id
            }
        )

        # --------------------------------------------------------
        # 5. Build Chroma records
        # --------------------------------------------------------

        ids = []
        metadatas = []

        for index, record in enumerate(
            chunk_records
        ):
            chunk = record["text"]
            page_number = record["page"]

            chunk_id = self._create_chunk_id(
                file_id,
                index,
                chunk,
            )

            ids.append(chunk_id)

            metadata = {
                "file_id": file_id,
                "source": path.name,
                "file_type": extension,
                "chunk_index": index,
            }

            # Add page only when page information exists.
            if page_number is not None:
                metadata["page"] = page_number

            metadatas.append(metadata)

        # --------------------------------------------------------
        # 6. Store in ChromaDB
        # --------------------------------------------------------

        self.collection.add(
            ids=ids,
            documents=chunks,
            embeddings=embeddings,
            metadatas=metadatas,
        )

        logger.info(
            "Successfully indexed %d chunks for %s",
            len(chunks),
            path.name,
        )

        return {
            "file_id": file_id,
            "filename": path.name,
            "file_type": extension,
            "chunk_count": len(chunks),
            "embedding_dimension": len(embeddings[0]),
            "status": "indexed",
        }

    # ============================================================
    # Document Extraction
    # ============================================================

    def _extract_document_pages(
        self,
        path: Path,
        extension: str,
    ) -> List[Tuple[Optional[int], str]]:
        """
        Extract document content while preserving page information.

        Returns:

            [
                (page_number, text),
                ...
            ]

        For PDF:
            page_number is the actual PDF page number.

        For DOCX/TXT:
            page_number is None.
        """

        if extension == ".pdf":
            return self._extract_pdf_pages(path)

        if extension == ".docx":
            text = self._extract_docx(path)

            return [
                (None, text)
            ]

        if extension == ".txt":
            text = self._extract_txt(path)

            return [
                (None, text)
            ]

        raise ValueError(
            f"Unsupported document type: {extension}"
        )

    # ============================================================
    # PDF Extraction
    # ============================================================

    def _extract_pdf_pages(
        self,
        path: Path,
    ) -> List[Tuple[int, str]]:
        """
        Extract text page-by-page from a PDF.

        Page numbering starts at 1.
        """

        reader = PdfReader(str(path))

        pages = []

        for page_number, page in enumerate(
            reader.pages,
            start=1,
        ):
            text = page.extract_text() or ""

            if text.strip():
                pages.append(
                    (
                        page_number,
                        text,
                    )
                )

        return pages

    # ============================================================
    # DOCX Extraction
    # ============================================================

    def _extract_docx(
        self,
        path: Path,
    ) -> str:
        """
        Extract paragraphs and tables from DOCX.
        """

        document = Document(str(path))

        parts: List[str] = []

        # --------------------------------------------------------
        # Paragraphs
        # --------------------------------------------------------

        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:
                parts.append(text)

        # --------------------------------------------------------
        # Tables
        # --------------------------------------------------------

        for table in document.tables:

            for row in table.rows:

                cells = [
                    cell.text.strip()
                    for cell in row.cells
                ]

                row_text = " | ".join(
                    cell
                    for cell in cells
                    if cell
                )

                if row_text:
                    parts.append(row_text)

        return "\n".join(parts)

    # ============================================================
    # TXT Extraction
    # ============================================================

    def _extract_txt(
        self,
        path: Path,
    ) -> str:
        """
        Extract UTF-8 text from a TXT file.
        """

        return path.read_text(
            encoding="utf-8",
            errors="replace",
        )

    # ============================================================
    # Text Processing
    # ============================================================

    def _normalize_text(
        self,
        text: str,
    ) -> str:
        """
        Normalize whitespace while preserving meaningful content.
        """

        lines = [
            line.strip()
            for line in text.splitlines()
        ]

        lines = [
            line
            for line in lines
            if line
        ]

        return "\n".join(lines)

    def _chunk_text(
        self,
        text: str,
        chunk_size: int,
        chunk_overlap: int,
    ) -> List[str]:
        """
        Split text into overlapping character chunks.

        Chunking is performed independently for each PDF page,
        allowing us to preserve the page number associated with
        every chunk.
        """

        if chunk_size <= 0:
            raise ValueError(
                "chunk_size must be greater than zero."
            )

        if chunk_overlap < 0:
            raise ValueError(
                "chunk_overlap cannot be negative."
            )

        if chunk_overlap >= chunk_size:
            raise ValueError(
                "chunk_overlap must be smaller than chunk_size."
            )

        chunks = []

        start = 0
        text_length = len(text)

        while start < text_length:

            end = min(
                start + chunk_size,
                text_length,
            )

            chunk = text[start:end].strip()

            if chunk:
                chunks.append(chunk)

            if end >= text_length:
                break

            start = end - chunk_overlap

        return chunks

    # ============================================================
    # IDs
    # ============================================================

    def _create_chunk_id(
        self,
        file_id: str,
        chunk_index: int,
        chunk: str,
    ) -> str:
        """
        Create a deterministic chunk ID.

        Re-ingesting the same file/chunk produces the same ID
        when the chunk position and content are unchanged.
        """

        raw = (
            f"{file_id}:"
            f"{chunk_index}:"
            f"{chunk}"
        )

        digest = hashlib.sha256(
            raw.encode("utf-8")
        ).hexdigest()

        return (
            f"{file_id}_"
            f"{chunk_index}_"
            f"{digest[:16]}"
        )


# ================================================================
# Singleton instance
# ================================================================

document_ingestion_service = DocumentIngestionService()